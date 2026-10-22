"""管理端考试成绩接口回归测试

覆盖：GET /exams/{id}/results（汇总口径/最高分聚合/未参考名单/权限/404）、
GET /exams/{id}/results/export（Excel 双 Sheet）、
GET /exams/{id}/results/{result_id}（答卷详情/未交卷 400/跨考试 404）、
GET /exams 列表附带的 assigned_count / submitted_user_count。
"""

import io
import uuid
import pytest
from openpyxl import load_workbook
from app.models.question import Question
from tests.conftest import create_question


def _q_single(db, stem="单选题", answer="A"):
    q = Question(
        id=str(uuid.uuid4()), type="single_choice", stem=stem,
        options={"A": "甲", "B": "乙", "C": "丙", "D": "丁"},
        answer={"correct": answer}, difficulty="medium",
    )
    db.add(q)
    db.commit()
    return q


def _mk_user(db, name):
    from app.models.user import User
    from app.services.auth_service import hash_password
    u = User(
        id=str(uuid.uuid4()),
        username=f"stu_{uuid.uuid4().hex[:8]}",
        password_hash=hash_password("Student123!"),
        name=name, role="student", status="active",
        department="测试部",
    )
    db.add(u)
    db.commit()
    return u


def _headers(client, user):
    res = client.post("/api/auth/login", json={
        "username": user.username, "password": "Student123!",
    })
    assert res.status_code == 200, res.text
    return {"Authorization": f"Bearer {res.json()['access_token']}"}


def _submit(client, headers, exam_id, qid, answer):
    res = client.post(f"/api/student/exams/{exam_id}/submit", headers=headers,
                      json={"answers": [{"question_id": qid, "answer": answer}]})
    assert res.status_code == 200, res.text
    return res.json()


@pytest.fixture()
def exam_env(client, admin_headers, course_id, db):
    """一场 10 分单选考试 + 5 名学员覆盖全部成绩形态：

    及格通过 / 未及格 / 两次作答取最高 / 作答中未交卷 / 未参考
    """
    q = _q_single(db)
    users = {name: _mk_user(db, name) for name in ("及格者", "未及格者", "重考者", "作答中", "未参考")}
    res = client.post("/api/exams/quick-create", headers=admin_headers, json={
        "title": "成绩测试考", "course_id": course_id,
        "content": {"source": "bank", "paper_type": "fixed",
                    "questions": [{"question_id": q.id, "score": 10}]},
        "pass_score": 6,
        "assign_user_ids": [u.id for u in users.values()],
    })
    assert res.status_code == 201, res.text
    exam_id = res.json()["id"]

    # 及格者：一次交卷 10 分
    h = _headers(client, users["及格者"])
    client.get(f"/api/student/exams/{exam_id}/start", headers=h)
    _submit(client, h, exam_id, q.id, "A")

    # 未及格者：交卷答错 → 0 分
    h = _headers(client, users["未及格者"])
    client.get(f"/api/student/exams/{exam_id}/start", headers=h)
    _submit(client, h, exam_id, q.id, "B")

    # 重考者：第一次答错、第二次答对 → 最高分 10、passed=True、交卷次数 2
    h = _headers(client, users["重考者"])
    client.get(f"/api/student/exams/{exam_id}/start", headers=h)
    _submit(client, h, exam_id, q.id, "B")
    client.get(f"/api/student/exams/{exam_id}/start", headers=h)
    _submit(client, h, exam_id, q.id, "A")

    # 作答中：start 后不交卷
    h = _headers(client, users["作答中"])
    client.get(f"/api/student/exams/{exam_id}/start", headers=h)

    return {"exam_id": exam_id, "q": q, "users": users}


# ========== 汇总 ==========

def test_results_404(client, admin_headers):
    res = client.get(f"/api/exams/{uuid.uuid4()}/results", headers=admin_headers)
    assert res.status_code == 404


def test_results_forbidden_for_student(client, student_headers):
    res = client.get("/api/exams/any-id/results", headers=student_headers)
    assert res.status_code == 403


def test_results_summary_counts(client, admin_headers, exam_env):
    res = client.get(f"/api/exams/{exam_env['exam_id']}/results", headers=admin_headers)
    assert res.status_code == 200, res.text
    data = res.json()
    s = data["summary"]
    assert s["assigned"] == 5
    assert s["submitted_users"] == 3          # 及格者/未及格者/重考者
    assert s["passed_users"] == 2             # 及格者/重考者（最高分口径）
    assert s["in_progress_users"] == 1        # 作答中
    assert s["absent_users"] == 1             # 未参考
    assert s["avg_score"] == pytest.approx(6.7)  # (10+0+10)/3
    assert data["exam"]["course_title"] == "测试课程"


def test_results_best_score_and_absent(client, admin_headers, exam_env):
    data = client.get(f"/api/exams/{exam_env['exam_id']}/results", headers=admin_headers).json()
    rows = {r["name"]: r for r in data["rows"]}

    retaker = rows["重考者"]
    assert retaker["best_score"] == 10
    assert retaker["passed"] is True
    assert retaker["attempt_count"] == 2
    assert len(retaker["attempts"]) == 2

    doing = rows["作答中"]
    assert doing["attempt_count"] == 0
    assert doing["best_score"] is None
    assert doing["in_progress"] is True
    assert all(a["submitted_at"] is None for a in doing["attempts"])

    assert [u["name"] for u in data["absent"]] == ["未参考"]
    # 排序：有分在前按降序，仅作答中（无分）的人排最后
    assert data["rows"][-1]["name"] == "作答中"


# ========== 答卷详情 ==========

def test_result_detail_ok(client, admin_headers, exam_env):
    data = client.get(f"/api/exams/{exam_env['exam_id']}/results", headers=admin_headers).json()
    row = next(r for r in data["rows"] if r["name"] == "及格者")
    result_id = row["attempts"][0]["id"]

    res = client.get(f"/api/exams/{exam_env['exam_id']}/results/{result_id}", headers=admin_headers)
    assert res.status_code == 200, res.text
    detail = res.json()
    assert detail["user_name"] == "及格者"
    assert detail["score"] == 10
    assert len(detail["questions"]) == 1
    q = detail["questions"][0]
    assert q["student_answer"] == "A"
    assert q["got_score"] == 10 and q["max_score"] == 10
    assert q["is_correct"] is True


def test_result_detail_in_progress_400(client, admin_headers, exam_env):
    data = client.get(f"/api/exams/{exam_env['exam_id']}/results", headers=admin_headers).json()
    row = next(r for r in data["rows"] if r["name"] == "作答中")
    res = client.get(f"/api/exams/{exam_env['exam_id']}/results/{row['attempts'][0]['id']}",
                     headers=admin_headers)
    assert res.status_code == 400
    assert "尚未交卷" in res.json()["detail"]


def test_result_detail_wrong_exam_404(client, admin_headers, exam_env, db):
    data = client.get(f"/api/exams/{exam_env['exam_id']}/results", headers=admin_headers).json()
    row = next(r for r in data["rows"] if r["name"] == "及格者")
    other_exam_id = str(uuid.uuid4())
    res = client.get(f"/api/exams/{other_exam_id}/results/{row['attempts'][0]['id']}",
                     headers=admin_headers)
    assert res.status_code == 404


# ========== Excel 导出 ==========

def test_results_export(client, admin_headers, exam_env):
    res = client.get(f"/api/exams/{exam_env['exam_id']}/results/export", headers=admin_headers)
    assert res.status_code == 200, res.text
    assert "spreadsheetml" in res.headers["content-type"]
    assert "filename*=UTF-8''" in res.headers.get("content-disposition", "")

    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["成绩", "未参考"]
    ws, ws2 = wb["成绩"], wb["未参考"]
    # 成绩：表头 + 4 行（有作答记录的人）
    assert ws.max_row == 5
    names = [ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)]
    assert "未参考" not in names and "重考者" in names
    # 未参考：表头 + 1 行
    assert ws2.max_row == 2
    assert ws2.cell(row=2, column=1).value == "未参考"


# ========== 列表附带统计 ==========

def test_exam_list_counts(client, admin_headers, exam_env):
    res = client.get("/api/exams", headers=admin_headers)
    assert res.status_code == 200
    item = next(e for e in res.json() if e["id"] == exam_env["exam_id"])
    assert item["assigned_count"] == 5
    assert item["submitted_user_count"] == 3
