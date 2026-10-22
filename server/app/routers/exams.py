"""考试管理路由"""

import io
import uuid
from datetime import datetime
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func
from sqlalchemy.orm import Session
from app.database import get_db
from app.models.user import User
from app.models.course import Course, CourseAssignment
from app.models.exam import Exam, ExamPaper, ExamResult
from app.models.paper import Paper
from app.models.question import Question
from app.schemas.exam import (
    ExamCreate, ExamUpdate, ExamOut, PaperCreate, PaperOut,
    ExamQuickCreate, ExamQuickCreateOut, PreviewDrawRequest,
)
from app.routers.auth import require_admin
from app.services.exam_service import (
    build_result_review, draw_random_questions, validate_random_rules,
    aggregate_exam_results,
)
from app.services.log_service import write_log
from app.logger import get_logger

logger = get_logger(__name__)

router = APIRouter(prefix="/api", tags=["考试管理"])


# ========== 考试 CRUD ==========

@router.get("/exams", response_model=list[ExamOut])
def list_exams(
    response: Response,
    course_id: str | None = Query(None),
    status: str | None = Query(None),
    page: int | None = Query(None, ge=1, description="页码（不传则返回全部，保持向后兼容）"),
    page_size: int | None = Query(None, ge=1, le=500, description="每页数量"),
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    """获取考试列表。

    新增可选分页：传入 page/page_size 时按页返回，并通过 `X-Total-Count` 响应头
    返回总数；不传时保持原有「返回全部」行为，兼容现有前端。
    """
    q = db.query(Exam)
    if course_id:
        q = q.filter(Exam.course_id == course_id)
    if status:
        q = q.filter(Exam.status == status)

    q = q.order_by(Exam.created_at.desc())
    if page and page_size:
        response.headers["X-Total-Count"] = str(q.count())
        exams = q.offset((page - 1) * page_size).limit(page_size).all()
    else:
        exams = q.all()

    # 附带「参考 X/Y」统计：两条分组查询，避免逐考试 N+1
    if not exams:
        return []
    exam_ids = [e.id for e in exams]
    course_ids = {e.course_id for e in exams}
    assigned_map = dict(
        db.query(CourseAssignment.course_id, func.count(func.distinct(CourseAssignment.user_id)))
        .filter(CourseAssignment.course_id.in_(course_ids))
        .group_by(CourseAssignment.course_id)
        .all()
    )
    submitted_map = dict(
        db.query(ExamResult.exam_id, func.count(func.distinct(ExamResult.user_id)))
        .filter(ExamResult.exam_id.in_(exam_ids), ExamResult.submitted_at.isnot(None))
        .group_by(ExamResult.exam_id)
        .all()
    )
    return [
        {
            **ExamOut.model_validate(e).model_dump(),
            "assigned_count": int(assigned_map.get(e.course_id, 0) or 0),
            "submitted_user_count": int(submitted_map.get(e.id, 0) or 0),
        }
        for e in exams
    ]


@router.post("/exams", response_model=ExamOut, status_code=201)
def create_exam(
    data: ExamCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    # 验证课程存在
    course = db.query(Course).filter(Course.id == data.course_id).first()
    if not course:
        raise HTTPException(status_code=404, detail="课程不存在")

    # 如果关联了试卷，读取试卷数据
    paper = None
    if data.paper_id:
        paper = db.query(Paper).filter(Paper.id == data.paper_id).first()
        if not paper:
            raise HTTPException(status_code=404, detail="试卷不存在")

    start_time = datetime.fromisoformat(data.start_time) if data.start_time else None
    end_time = datetime.fromisoformat(data.end_time) if data.end_time else None

    exam = Exam(
        id=str(uuid.uuid4()),
        course_id=data.course_id,
        title=data.title,
        paper_type=data.paper_type,
        paper_id=data.paper_id,
        duration=data.duration,
        total_score=data.total_score,
        pass_score=data.pass_score,
        max_attempts=data.max_attempts,
        start_time=start_time,
        end_time=end_time,
        status=data.status,
        created_by=current_user.id,
    )
    db.add(exam)
    db.flush()  # 获取 exam.id，但暂不提交

    # 从试卷复制组卷内容到 exam_papers
    if paper:
        exam_paper = ExamPaper(
            id=str(uuid.uuid4()),
            exam_id=exam.id,
            questions=paper.mandatory_questions,
            random_rules=paper.random_rules,
        )
        db.add(exam_paper)

    # 将考试关联到课程，学员端通过 course.exam_id 看到考试。
    # 数据模型是课程↔考试 1:1；若该课程原先绑定的是另一个考试，把旧考试自动置为
    # 「已结束」，避免它继续以「已发布」状态存在却无法被学员访问（原来的静默覆盖
    # 会让旧考试变成没人知道的孤儿）。
    _bind_exam_to_course(db, course, exam)

    db.commit()
    db.refresh(exam)
    write_log(db=db, operator=current_user.name, operator_id=current_user.id, category="exam", action="create", resource_type="考试", resource_name=exam.title if hasattr(exam, 'title') else "", description=f"{current_user.name} 创建了考试「{exam.title}」", level="info")
    return exam


@router.get("/exams/{exam_id}", response_model=ExamOut)
def get_exam(
    exam_id: str,
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="考试不存在")
    return exam


@router.put("/exams/{exam_id}", response_model=ExamOut)
def update_exam(
    exam_id: str,
    data: ExamUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="考试不存在")
    update_data = data.model_dump(exclude_unset=True)
    # 处理时间字段
    for field in ("start_time", "end_time"):
        if field in update_data and update_data[field] is not None:
            update_data[field] = datetime.fromisoformat(update_data[field])
    for field, value in update_data.items():
        setattr(exam, field, value)
    db.commit()
    db.refresh(exam)
    write_log(db=db, operator=current_user.name, operator_id=current_user.id, category="exam", action="update", resource_type="考试", resource_name=exam.title if hasattr(exam, 'title') else "", description=f"{current_user.name} 编辑了考试「{exam.title}」", level="info")
    return exam


@router.delete("/exams/{exam_id}")
def delete_exam(
    exam_id: str,
    confirm: bool = Query(False, description="必须显式传 true 才执行删除"),
    force: bool = Query(False, description="已有学员作答记录时，需再传 force=true 才会连历史成绩一并删除"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    """删除考试。

    修复（防误删 + 保护历史）：
    - 必须显式 `confirm=true`
    - 若已有学员作答记录，默认**拒绝删除**（返回 409），
      请改为把考试状态置为 `ended`（结束）以保留成绩；确实要连成绩一起删，再传 `force=true`
    """
    if not confirm:
        raise HTTPException(status_code=400, detail="删除考试需要显式确认（confirm=true）")

    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="考试不存在")

    result_count = db.query(ExamResult).filter(ExamResult.exam_id == exam_id).count()
    if result_count and not force:
        raise HTTPException(
            status_code=409,
            detail=(
                f"该考试已有 {result_count} 条学员作答记录，为避免误删历史成绩已阻止删除。"
                "如只想停止考试，请把状态改为「已结束」；如确认要连成绩一并删除，请传 force=true。"
            ),
        )

    title = exam.title
    # 清理课程关联
    db.query(Course).filter(Course.exam_id == exam_id).update(
        {"exam_id": None}, synchronize_session=False
    )

    db.query(ExamResult).filter(ExamResult.exam_id == exam_id).delete()
    db.query(ExamPaper).filter(ExamPaper.exam_id == exam_id).delete()
    db.delete(exam)
    db.commit()
    write_log(db=db, operator=current_user.name, operator_id=current_user.id, category="exam", action="delete", resource_type="考试", resource_name=title, description=f"{current_user.name} 删除了考试「{title}」", level="warning")
    return {"ok": True}


# ========== 组卷 ==========

@router.get("/exams/{exam_id}/paper", response_model=PaperOut)
def get_paper(
    exam_id: str,
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    paper = db.query(ExamPaper).filter(ExamPaper.exam_id == exam_id).first()
    if not paper:
        return PaperOut(id="", exam_id=exam_id, questions=None, random_rules=None)
    return paper


@router.post("/exams/{exam_id}/paper", response_model=PaperOut)
def save_paper(
    exam_id: str,
    data: PaperCreate,
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="考试不存在")

    # 删除旧试卷
    db.query(ExamPaper).filter(ExamPaper.exam_id == exam_id).delete()

    questions = [q.model_dump() for q in data.questions] if data.questions else None
    random_rules = [r.model_dump() for r in data.random_rules] if data.random_rules else None

    paper = ExamPaper(
        id=str(uuid.uuid4()),
        exam_id=exam_id,
        questions=questions,
        random_rules=random_rules,
    )
    db.add(paper)
    db.commit()
    db.refresh(paper)
    return paper


# ========== 成绩查看（管理端） ==========

def _get_exam_or_404(db: Session, exam_id: str) -> Exam:
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="考试不存在")
    return exam


@router.get("/exams/{exam_id}/results")
def get_exam_results(
    exam_id: str,
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    """考试成绩汇总：应考/已交卷/通过/作答中/未参考 + 按人分组（多次作答取最高分）。

    口径见 app.services.exam_service.aggregate_exam_results 文档。
    """
    exam = _get_exam_or_404(db, exam_id)
    return aggregate_exam_results(db, exam)


@router.get("/exams/{exam_id}/results/export")
def export_exam_results(
    exam_id: str,
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    """导出成绩 Excel：Sheet1 成绩（一人一行，含各次明细）、Sheet2 未参考名单。

    注意：本路由必须注册在 `/exams/{exam_id}/results/{result_id}` 之前，
    否则 "export" 会被当作 result_id 捕获（同 questions.py 的历史教训）。
    """
    exam = _get_exam_or_404(db, exam_id)
    data = aggregate_exam_results(db, exam)

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="EFEFEF")
    center = Alignment(horizontal="center")

    def style_header(ws):
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

    wb = Workbook()
    ws = wb.active
    ws.title = "成绩"
    ws.append(["姓名", "工号", "部门", "最高分", "是否通过", "交卷次数", "状态", "各次成绩", "最近交卷时间"])
    for row in data["rows"]:
        submitted = sorted(
            (a for a in row["attempts"] if a["submitted_at"]),
            key=lambda a: a["attempt_number"],
        )
        detail = "；".join(
            f"第{a['attempt_number']}次 {a['score']}分({'通过' if a['passed'] else '未通过'})"
            for a in submitted
        )
        if row["attempt_count"] > 0:
            status = "已通过" if row["passed"] else "未通过"
            if row["in_progress"]:
                status += "（另有作答中）"
        else:
            status = "作答中"
        last_submitted = max((a["submitted_at"] for a in submitted), default="") or ""
        ws.append([
            row["name"],
            row["employee_id"] or "",
            row["department"] or "",
            row["best_score"] if row["best_score"] is not None else "",
            "通过" if row["passed"] else ("未通过" if row["attempt_count"] > 0 else ""),
            row["attempt_count"],
            status,
            detail,
            last_submitted.replace("T", " "),
        ])
    style_header(ws)
    for col, width in zip("ABCDEFGHI", (14, 12, 24, 8, 10, 10, 18, 60, 20)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("未参考")
    ws2.append(["姓名", "工号", "部门"])
    for u in data["absent"]:
        ws2.append([u["name"], u["employee_id"] or "", u["department"] or ""])
    style_header(ws2)
    for col, width in zip("ABC", (14, 12, 30)):
        ws2.column_dimensions[col].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"{exam.title}-成绩.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f"attachment; filename=exam-results.xlsx; filename*=UTF-8''{quote(filename)}"
            ),
        },
    )


@router.get("/exams/{exam_id}/results/{result_id}")
def get_exam_result_detail(
    exam_id: str,
    result_id: str,
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    """管理端答卷详情：按题目快照逐题比对（与学员端回顾同一套 build_result_review）。"""
    exam = _get_exam_or_404(db, exam_id)

    result = db.query(ExamResult).filter(ExamResult.id == result_id).first()
    if not result or result.exam_id != exam_id:
        raise HTTPException(status_code=404, detail="考试记录不存在")
    if not result.submitted_at:
        raise HTTPException(status_code=400, detail="该考生尚未交卷，无法查看答卷")
    if not result.questions:
        raise HTTPException(status_code=400, detail="该考试记录缺少题目快照，无法查看答卷")

    review = build_result_review(db, exam, result)
    taker = db.query(User).filter(User.id == result.user_id).first()
    review["user_name"] = taker.name if taker else "(已删除用户)"
    review["employee_id"] = taker.employee_id if taker else None
    review["department"] = taker.department if taker else None
    return review


# ========== 一键建考 ==========

@router.post("/exams/preview-draw")
def preview_draw(
    data: PreviewDrawRequest,
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    """服务端随机抽题预览：与学员端开考用同一套抽题逻辑，所见即所得。

    题库存量不足时不阻断（按现有存量抽出），通过 shortages 字段返回缺题说明，
    由向导前端提示管理员调整规则。
    """
    rules = [r.model_dump() for r in data.random_rules]
    problems = validate_random_rules(db, rules, data.source_banks)
    questions = draw_random_questions(db, rules, data.source_banks)
    return {
        "questions": questions,
        "total_score": sum(q["score"] for q in questions),
        "shortages": problems,
    }


def _bind_exam_to_course(db: Session, course: Course, exam: Exam) -> str | None:
    """把考试绑定到课程（课程↔考试 1:1），返回被自动置为「已结束」的旧考试名称。"""
    previous_exam_id = course.exam_id
    replaced_title = None
    if previous_exam_id and previous_exam_id != exam.id:
        previous = db.query(Exam).filter(Exam.id == previous_exam_id).first()
        if previous and previous.status != "ended":
            previous.status = "ended"
            replaced_title = previous.title
            logger.warning(
                "课程「%s」的考试绑定已由「%s」替换为「%s」，旧考试自动置为已结束",
                course.title, previous.title, exam.title,
            )
    course.exam_id = exam.id
    return replaced_title


@router.post("/exams/quick-create", response_model=ExamQuickCreateOut, status_code=201)
def quick_create_exam(
    data: ExamQuickCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    """一键建考：选题/复用试卷 + 建考试 + 组卷 + 绑课程 + （默认）发布，一个事务完成。

    智能默认：
    - total_score 留空 → 按试卷内容自动累加；固定组卷且未设分值时，留空的
      total_score 作为平摊目标（默认 100 分），整数平摊、余数补给前面的题
    - pass_score 留空 → round(总分 × 60%)
    - publish=true（默认）→ 创建即发布，省掉「发布试卷→发布考试」两步
    """
    course = db.query(Course).filter(Course.id == data.course_id).first()
    if not course:
        raise HTTPException(status_code=404, detail="课程不存在")

    content = data.content
    questions_cfg: list[dict] | None = None
    rules_cfg: list[dict] | None = None
    banks_cfg: list[str] | None = None
    paper_type = content.paper_type
    ref_paper_id: str | None = None

    if content.source == "paper":
        # ── 复用已有试卷：与旧 POST /exams 的拷贝行为一致 ──
        if not content.paper_id:
            raise HTTPException(status_code=400, detail="复用试卷时必须指定 paper_id")
        src_paper = db.query(Paper).filter(Paper.id == content.paper_id).first()
        if not src_paper:
            raise HTTPException(status_code=404, detail="试卷不存在")
        if src_paper.status != "published":
            raise HTTPException(status_code=400, detail="所选试卷尚未发布，请先在试卷管理中发布")
        paper_type = src_paper.paper_type
        ref_paper_id = src_paper.id
        if paper_type == "fixed":
            questions_cfg = [
                {"question_id": q["question_id"], "score": q.get("score", 0)}
                for q in (src_paper.mandatory_questions or [])
            ]
            if not questions_cfg:
                raise HTTPException(status_code=400, detail="所选试卷尚未配置任何题目")
            if any(q["score"] <= 0 for q in questions_cfg):
                raise HTTPException(status_code=400, detail="所选试卷存在分值未设置的题目，请先在试卷管理中补充分值")
            computed_total = sum(q["score"] for q in questions_cfg)
        else:
            rules_cfg = src_paper.random_rules or []
            if not rules_cfg:
                raise HTTPException(status_code=400, detail="所选试卷尚未配置抽题规则")
            banks_cfg = src_paper.source_banks
            problems = validate_random_rules(db, rules_cfg, banks_cfg)
            if problems:
                raise HTTPException(status_code=400, detail="题库缺题，无法创建：" + "；".join(problems))
            computed_total = sum(int(r["count"]) * int(r["score"]) for r in rules_cfg)
    elif content.source == "bank":
        # ── 直接从题库选题 ──
        if paper_type == "fixed":
            picked = content.questions or []
            if not picked:
                raise HTTPException(status_code=400, detail="请至少选择一道题目")
            ids = [q.question_id for q in picked]
            if len(set(ids)) != len(ids):
                raise HTTPException(status_code=400, detail="选题列表中存在重复题目")
            found = {
                q.id for q in db.query(Question).filter(Question.id.in_(ids)).all()
            }
            missing = [i for i in ids if i not in found]
            if missing:
                raise HTTPException(status_code=400, detail=f"有 {len(missing)} 道题目不存在或已被删除，请重新选题")
            scored = [q.score or 0 for q in picked]
            if all(s > 0 for s in scored):
                questions_cfg = [{"question_id": q.question_id, "score": s} for q, s in zip(picked, scored)]
                computed_total = sum(scored)
            else:
                # 未逐题设分值 → 按目标总分整数平摊，余数补给前面的题
                target_total = data.total_score or 100
                n = len(picked)
                base, rem = divmod(target_total, n)
                if base <= 0:
                    raise HTTPException(
                        status_code=400,
                        detail=f"目标总分 {target_total} 不足以平摊给 {n} 道题，请将总分提高到至少 {n} 分或减少题量",
                    )
                questions_cfg = [
                    {"question_id": q.question_id, "score": base + (1 if i < rem else 0)}
                    for i, q in enumerate(picked)
                ]
                computed_total = sum(q["score"] for q in questions_cfg)
        elif paper_type == "random":
            rules_cfg = [r.model_dump() for r in (content.random_rules or [])]
            rules_cfg = [r for r in rules_cfg if int(r.get("count") or 0) > 0]
            if not rules_cfg:
                raise HTTPException(status_code=400, detail="请至少配置一条抽题数量大于 0 的随机规则")
            if any(int(r["score"]) <= 0 for r in rules_cfg):
                raise HTTPException(status_code=400, detail="随机规则每题分值必须大于 0")
            problems = validate_random_rules(db, rules_cfg, content.source_banks)
            if problems:
                raise HTTPException(status_code=400, detail="题库缺题，无法创建：" + "；".join(problems))
            banks_cfg = content.source_banks
            computed_total = sum(int(r["count"]) * int(r["score"]) for r in rules_cfg)
        else:
            raise HTTPException(status_code=400, detail="组卷方式无效，应为 fixed / random")
    else:
        raise HTTPException(status_code=400, detail="内容来源无效，应为 bank / paper")

    total_score = data.total_score or computed_total
    pass_score = data.pass_score if data.pass_score is not None else max(1, round(total_score * 0.6))
    if pass_score > total_score:
        raise HTTPException(status_code=400, detail=f"及格分（{pass_score}）不能大于总分（{total_score}）")

    start_time = datetime.fromisoformat(data.start_time) if data.start_time else None
    end_time = datetime.fromisoformat(data.end_time) if data.end_time else None
    if start_time and end_time and start_time >= end_time:
        raise HTTPException(status_code=400, detail="开放时间必须早于截止时间")

    # 建考即分配学员：先校验后插入，与建考同一事务；校验不过直接 400，不留任何残留
    assigned_new = 0
    if data.assign_user_ids:
        assign_ids = list(dict.fromkeys(data.assign_user_ids))  # 去重保序
        found = {u.id for u in db.query(User).filter(User.id.in_(assign_ids)).all()}
        missing = [i for i in assign_ids if i not in found]
        if missing:
            raise HTTPException(status_code=400, detail=f"有 {len(missing)} 名学员不存在或已被删除，请重新选择")
        existing = {
            a.user_id for a in db.query(CourseAssignment).filter(
                CourseAssignment.course_id == course.id,
                CourseAssignment.user_id.in_(assign_ids),
            ).all()
        }
        for uid in assign_ids:
            if uid not in existing:
                db.add(CourseAssignment(
                    id=str(uuid.uuid4()), course_id=course.id, user_id=uid,
                    assigned_by=current_user.id,
                ))
                assigned_new += 1

    exam = Exam(
        id=str(uuid.uuid4()),
        course_id=course.id,
        title=data.title,
        paper_type=paper_type,
        paper_id=ref_paper_id,
        duration=data.duration,
        total_score=total_score,
        pass_score=pass_score,
        max_attempts=data.max_attempts,
        start_time=start_time,
        end_time=end_time,
        status="published" if data.publish else "draft",
        created_by=current_user.id,
    )
    db.add(exam)
    db.flush()

    db.add(ExamPaper(
        id=str(uuid.uuid4()),
        exam_id=exam.id,
        questions=questions_cfg,
        random_rules=rules_cfg,
        source_banks=banks_cfg,
    ))

    replaced_title = _bind_exam_to_course(db, course, exam)

    saved_paper_id = None
    if data.save_paper_as:
        tpl = Paper(
            id=str(uuid.uuid4()),
            title=data.save_paper_as,
            category_id=content.category_id,
            paper_type=paper_type,
            mandatory_questions=questions_cfg,
            random_rules=rules_cfg,
            source_banks=banks_cfg,
            total_score=computed_total,
            status="published",
            created_by=current_user.id,
        )
        db.add(tpl)
        saved_paper_id = tpl.id

    db.commit()
    db.refresh(exam)
    write_log(
        db=db, operator=current_user.name, operator_id=current_user.id, category="exam",
        action="create", resource_type="考试", resource_name=exam.title,
        description=(
            f"{current_user.name} 一键创建了考试「{exam.title}」"
            f"（{'已发布' if exam.status == 'published' else '草稿'}，"
            f"{'复用试卷' if content.source == 'paper' else '题库选题'}·{paper_type}）"
        ),
        level="info",
    )
    out = ExamOut.model_validate(exam).model_dump()
    out["replaced_exam_title"] = replaced_title
    out["saved_paper_id"] = saved_paper_id
    out["assigned_new"] = assigned_new
    return out