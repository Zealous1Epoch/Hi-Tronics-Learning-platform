"""参考资料文档文本解析（AI 出题上传文件 → 纯文本）

- 只解析不落库：结果由前端放入生成请求的 reference_text，prompt 端三引号定界防注入
- PDF / Word(.docx) / 纯文本走服务端本地库解析；图片走视觉 LLM OCR（需 LLM_VISION_MODEL）
- pypdf / python-docx 均为函数内 import：未安装时其他 AI 功能不受影响，仅在解析时报错
"""

import base64
import io

from fastapi import HTTPException

from app.config import settings

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 20000

PDF_EXTS = {"pdf"}
DOC_EXTS = {"docx"}
TEXT_EXTS = {"txt", "md"}
IMAGE_EXTS = {"png", "jpg", "jpeg", "webp", "bmp", "gif"}
ALLOWED_EXTS = PDF_EXTS | DOC_EXTS | TEXT_EXTS | IMAGE_EXTS

_IMAGE_MIME = {
    "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
    "webp": "image/webp", "bmp": "image/bmp", "gif": "image/gif",
}

_OCR_PROMPT = (
    "提取这张图片中的全部文字，按自然阅读顺序输出。"
    "只输出提取的文本本身，不要添加任何解释、总结或格式标记；"
    "若图片中没有文字，仅输出：（图片无文字）"
)


def ext_of(filename: str) -> str:
    base = (filename or "").lower()
    return base.rsplit(".", 1)[1] if "." in base else ""


def _extract_pdf(blob: bytes) -> tuple[str, str | None]:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise HTTPException(status_code=501, detail="服务端未安装 pypdf，无法解析 PDF")
    try:
        reader = PdfReader(io.BytesIO(blob))
        pages = [(p.extract_text() or "") for p in reader.pages]
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"PDF 解析失败（{type(e).__name__}），请确认文件未损坏且未加密",
        )
    text = "\n\n".join(t.strip() for t in pages if t.strip())
    note = None
    if not text:
        note = "未在 PDF 中提取到可选文字（可能是扫描件），可将页面截图为图片上传识别"
    return text, note


def _extract_docx(blob: bytes) -> tuple[str, str | None]:
    try:
        from docx import Document
    except ImportError:
        raise HTTPException(status_code=501, detail="服务端未安装 python-docx，无法解析 Word")
    try:
        doc = Document(io.BytesIO(blob))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        # 表格文本统一追加在段落后（保留内容优先于还原版式）
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append("　".join(cells))
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Word 解析失败（{type(e).__name__}），仅支持 .docx，旧版 .doc 请先另存为 .docx",
        )
    return "\n".join(parts), None


def _extract_textfile(blob: bytes) -> tuple[str, str | None]:
    for enc in ("utf-8", "gb18030"):
        try:
            return blob.decode(enc), None
        except UnicodeDecodeError:
            continue
    return blob.decode("utf-8", errors="ignore"), "文本存在无法识别的编码字符，已跳过"


def _extract_image(blob: bytes, ext: str) -> tuple[str, str | None]:
    if not settings.LLM_VISION_MODEL:
        raise HTTPException(
            status_code=400,
            detail="图片文字识别需在服务端配置 LLM_VISION_MODEL（支持视觉的模型），当前未配置",
        )
    from app.llm.provider import chat_completion_vision

    try:
        text = chat_completion_vision(
            _OCR_PROMPT, base64.b64encode(blob).decode(), _IMAGE_MIME.get(ext, "image/png")
        )
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"视觉模型识别图片失败（{type(e).__name__}）")
    return (text or "").strip(), None


def parse_document(filename: str, blob: bytes) -> dict:
    """按扩展名分发解析，返回 {name, ext, chars, truncated, text, note}。"""
    if len(blob) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="文件过大，单个上限 20MB")
    ext = ext_of(filename)
    if ext not in ALLOWED_EXTS:
        raise HTTPException(
            status_code=400,
            detail=f"不支持的文件类型 .{ext or '?'}（支持 PDF / Word(.docx) / 图片 / txt / md）",
        )
    if not blob:
        raise HTTPException(status_code=400, detail="文件内容为空")

    if ext in PDF_EXTS:
        text, note = _extract_pdf(blob)
    elif ext in DOC_EXTS:
        text, note = _extract_docx(blob)
    elif ext in IMAGE_EXTS:
        text, note = _extract_image(blob, ext)
    else:
        text, note = _extract_textfile(blob)

    text = text.strip()
    if not text and not note:
        note = "未识别到有效文字内容"
    truncated = len(text) > MAX_TEXT_CHARS
    if truncated:
        text = text[:MAX_TEXT_CHARS]
    return {
        "name": (filename or "file").rsplit("/", 1)[-1],
        "ext": ext,
        "chars": len(text),
        "truncated": truncated,
        "text": text,
        "note": note,
    }
