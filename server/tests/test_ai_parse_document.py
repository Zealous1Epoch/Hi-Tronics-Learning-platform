"""POST /api/ai/parse-document 参考资料解析测试（零网络，视觉路 monkeypatch）"""

import io

from app.config import settings
# 显式注册 ORM 模型：conftest 的 db fixture 在 create_all 前需要 Base.metadata 已含 users 表
import app.models.user  # noqa: F401
from app.services.doc_parse_service import MAX_UPLOAD_BYTES


def _up(client, headers, name, blob, ctype="application/octet-stream"):
    return client.post(
        "/api/ai/parse-document",
        files={"file": (name, io.BytesIO(blob), ctype)},
        headers=headers,
    )


class TestParseDocument:
    def test_txt_utf8(self, client, admin_headers):
        res = _up(client, admin_headers, "车间制度.txt", "安全第一、预防为主".encode("utf-8"))
        assert res.status_code == 200
        body = res.json()
        assert body["text"] == "安全第一、预防为主"
        assert body["chars"] == 9 and body["ext"] == "txt"

    def test_txt_gb18030_fallback(self, client, admin_headers):
        res = _up(client, admin_headers, "a.txt", "三级安全教育".encode("gb18030"))
        assert res.json()["text"] == "三级安全教育"

    def test_docx_paragraphs_and_tables(self, client, admin_headers):
        from docx import Document

        doc = Document()
        doc.add_paragraph("新员工入职须完成三级安全教育。")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "厂级"
        table.rows[0].cells[1].text = "安全法规"
        buf = io.BytesIO()
        doc.save(buf)

        res = _up(client, admin_headers, "课件.docx", buf.getvalue())
        assert res.status_code == 200
        text = res.json()["text"]
        assert "三级安全教育" in text and "厂级" in text and "安全法规" in text

    def test_broken_pdf_400(self, client, admin_headers):
        res = _up(client, admin_headers, "坏文件.pdf", b"definitely not a pdf")
        assert res.status_code == 400
        assert "PDF 解析失败" in res.json()["detail"]

    def test_scanned_pdf_returns_note(self, client, admin_headers, monkeypatch):
        class EmptyReader:
            pages = []

        monkeypatch.setattr("pypdf.PdfReader", lambda f: EmptyReader())
        res = _up(client, admin_headers, "扫描件.pdf", b"%PDF-fake")
        assert res.status_code == 200
        body = res.json()
        assert body["text"] == ""
        assert "扫描件" in body["note"]

    def test_image_requires_vision_model(self, client, admin_headers, monkeypatch):
        monkeypatch.setattr(settings, "LLM_VISION_MODEL", "")
        res = _up(client, admin_headers, "截图.png", b"\x89PNG fake")
        assert res.status_code == 400
        assert "LLM_VISION_MODEL" in res.json()["detail"]

    def test_image_via_vision_model(self, client, admin_headers, monkeypatch):
        captured = {}

        def fake_vision(prompt, image_b64, mime):
            captured["mime"] = mime
            return "灭火器压力表分三色：红黄绿"

        monkeypatch.setattr(settings, "LLM_VISION_MODEL", "qwen-vl-plus")
        monkeypatch.setattr("app.llm.provider.chat_completion_vision", fake_vision)
        res = _up(client, admin_headers, "截图.jpg", b"fakejpegbytes")
        assert res.status_code == 200
        assert res.json()["text"] == "灭火器压力表分三色：红黄绿"
        assert captured["mime"] == "image/jpeg"

    def test_unsupported_ext_400(self, client, admin_headers):
        res = _up(client, admin_headers, "virus.exe", b"MZ...")
        assert res.status_code == 400
        assert "不支持" in res.json()["detail"]

    def test_oversize_413(self, client, admin_headers):
        res = _up(client, admin_headers, "big.txt", b"a" * (MAX_UPLOAD_BYTES + 1))
        assert res.status_code == 413

    def test_text_truncated_at_cap(self, client, admin_headers):
        res = _up(client, admin_headers, "long.txt", ("知" * 25000).encode("utf-8"))
        body = res.json()
        assert body["chars"] == 20000 and body["truncated"] is True

    def test_requires_admin(self, client, student_headers):
        res = _up(client, student_headers, "a.txt", b"x")
        assert res.status_code == 403
