"""题库管理路由"""

import uuid
import io
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel
from app.database import get_db
from app.models.user import User
from app.models.question import Question
from app.models.paper import Paper
from app.models.exam import ExamPaper
from app.schemas.question import QuestionCreate, QuestionUpdate, QuestionOut, PaginatedResponse
from app.schemas.ai_question import BulkCreateRequest
from app.routers.auth import require_admin
from app.services.ai_question_service import normalize_stem, validate_item
from app.services.log_service import write_log

# AI 出题入库统一打的标签
AI_TAG = "AI生成"

router = APIRouter(prefix="/api", tags=["题库管理"])


@router.get("/questions", response_model=PaginatedResponse[QuestionOut])
def list_questions(
    type: str | None = Query(None, description="题型筛选"),
    difficulty: str | None = Query(None, description="难度筛选"),
    category_id: str | None = Query(None, description="分类筛选"),
    search: str | None = Query(None, description="搜索题干"),
    ids: str | None = Query(None, description="逗号分隔的题目 ID，传入时仅返回这些题且不分页（复制考试预填选题用）"),
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(50, ge=1, le=500, description="每页数量"),
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    q = db.query(Question)
    if ids:
        id_list = [i.strip() for i in ids.split(",") if i.strip()][:500]
        q = q.filter(Question.id.in_(id_list))
        items = q.all()
        return {"items": items, "total": len(items), "page": 1, "page_size": len(items) or 1}
    if type:
        q = q.filter(Question.type == type)
    if difficulty:
        q = q.filter(Question.difficulty == difficulty)
    if category_id:
        q = q.filter(Question.category_id == category_id)
    if search:
        q = q.filter(Question.stem.contains(search))
    total = q.count()
    items = q.order_by(Question.created_at.desc()).offset((page - 1) * page_size).limit(page_size).all()
    return {"items": items, "total": total, "page": page, "page_size": page_size}


@router.post("/questions", response_model=QuestionOut, status_code=201)
def create_question(
    data: QuestionCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    # 校验选择题必须有 options
    if data.type in ("single_choice", "multi_choice") and not data.options:
        raise HTTPException(status_code=400, detail="选择题必须提供选项")
    question = Question(
        id=str(uuid.uuid4()),
        type=data.type,
        stem=data.stem,
        options=data.options,
        answer=data.answer,
        keywords=data.keywords,
        difficulty=data.difficulty,
        tags=data.tags,
        section=data.section,
        category_id=data.category_id,
        created_by=current_user.id,
    )
    db.add(question)
    db.commit()
    db.refresh(question)
    stem = data.stem
    write_log(db=db, operator=current_user.name, operator_id=current_user.id, category="question", action="create", resource_type="试题", resource_name=stem[:50] if stem else "", description=f"{current_user.name} 创建了试题「{stem[:50]}」", level="info")
    return question


# ========== 批量操作 ==========

class BatchRequest(BaseModel):
    ids: list[str]
    action: str  # update_difficulty / add_tags / delete
    value: str | list[str] | None = None


@router.post("/questions/batch")
def batch_operation(
    data: BatchRequest,
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    """批量操作：修改难度、追加标签、删除"""
    questions = db.query(Question).filter(Question.id.in_(data.ids)).all()
    if not questions:
        raise HTTPException(status_code=404, detail="未找到试题")

    if data.action == "update_difficulty":
        if not data.value or data.value not in ("easy", "medium", "hard"):
            raise HTTPException(status_code=400, detail="无效的难度值")
        for q in questions:
            q.difficulty = data.value

    elif data.action == "add_tags":
        if not data.value:
            raise HTTPException(status_code=400, detail="标签不能为空")
        new_tags = data.value if isinstance(data.value, list) else [data.value]
        for q in questions:
            existing = set(q.tags or [])
            existing.update(new_tags)
            q.tags = list(existing)

    elif data.action == "delete":
        refs = _check_references(db, data.ids)
        ref_ids = {r["question_id"] for r in refs}
        deleted = 0
        skipped = 0
        for q in questions:
            if q.id in ref_ids:
                skipped += 1
                continue
            db.delete(q)
            deleted += 1
        db.commit()
        return {"ok": True, "deleted": deleted, "skipped": skipped, "refs": refs}

    else:
        raise HTTPException(status_code=400, detail=f"未知操作: {data.action}")

    db.commit()
    return {"ok": True, "affected": len(questions)}


# ========== 导入导出 ==========

TYPE_MAP = {
    "单选题": "single_choice", "single_choice": "单选题",
    "多选题": "multi_choice", "multi_choice": "多选题",
    "判断题": "true_false", "true_false": "判断题",
    "填空题": "fill_blank", "fill_blank": "填空题",
    "问答题": "essay", "essay": "问答题",
}


@router.get("/questions/export")
def export_questions(
    category_id: str | None = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    """导出题库为 Excel"""
    query = db.query(Question)
    if category_id:
        query = query.filter(Question.category_id == category_id)
    questions = query.order_by(Question.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "题库"

    # 表头
    headers = ["题型", "题干", "选项A", "选项B", "选项C", "选项D", "正确答案", "关键词", "标签", "难度"]
    ws.append(headers)

    for q in questions:
        options = q.options or {}
        answer = q.answer or {}
        row = [
            TYPE_MAP.get(q.type, q.type),
            q.stem,
            options.get("A", ""),
            options.get("B", ""),
            options.get("C", ""),
            options.get("D", ""),
            _format_answer(q.type, answer),
            ",".join(q.keywords) if q.keywords else "",
            ",".join(q.tags) if q.tags else "",
            q.difficulty,
        ]
        ws.append(row)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=questions.xlsx"},
    )


@router.get("/questions/template")
def download_template(_: User = Depends(require_admin)):
    """下载导入模板（修复：原实现缺少鉴权，任何匿名用户均可下载）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "导入模板"

    # 样式定义
    header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="C0C0C0"),
        right=Side(style="thin", color="C0C0C0"),
        top=Side(style="thin", color="C0C0C0"),
        bottom=Side(style="thin", color="C0C0C0"),
    )
    sample_font = Font(name="微软雅黑", size=10)
    tip_font = Font(name="微软雅黑", size=9, color="888888", italic=True)

    # 表头
    headers = ["题型", "题干", "选项A", "选项B", "选项C", "选项D", "正确答案", "关键词", "标签", "难度"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 示例行
    samples = [
        ["单选题", "Python 中列表和元组的区别是什么？", "可变", "不可变", "有序", "无序", "A", "", "Python,基础", "easy"],
        ["多选题", "以下哪些是 Python 不可变类型？", "字符串", "列表", "元组", "字典", "A,C", "", "Python,数据类型", "medium"],
        ["判断题", "Python 是解释型语言", "", "", "", "", "true", "", "Python", "easy"],
        ["填空题", "Python 中定义函数使用___关键字", "", "", "", "", "def", "", "Python,函数", "easy"],
        ["问答题", "简述面向对象编程的三大特性", "", "", "", "", "", "封装,继承,多态", "OOP", "hard"],
    ]
    for row_idx, sample in enumerate(samples, start=2):
        for col_idx, val in enumerate(sample, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = sample_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = thin_border

    # 列宽
    col_widths = [10, 40, 12, 12, 12, 12, 12, 20, 16, 10]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # 冻结首行
    ws.freeze_panes = "A2"

    # 数据验证：题型下拉
    type_dv = DataValidation(
        type="list",
        formula1='"单选题,多选题,判断题,填空题,问答题"',
        allow_blank=False,
    )
    type_dv.error = "请从下拉列表中选择题型"
    type_dv.errorTitle = "输入错误"
    type_dv.prompt = "请选择题型"
    type_dv.promptTitle = "题型"
    ws.add_data_validation(type_dv)
    type_dv.add(f"A2:A1000")

    # 数据验证：难度下拉
    diff_dv = DataValidation(
        type="list",
        formula1='"easy,medium,hard"',
        allow_blank=False,
    )
    diff_dv.error = "请从下拉列表中选择难度"
    diff_dv.errorTitle = "输入错误"
    diff_dv.prompt = "请选择难度"
    diff_dv.promptTitle = "难度"
    ws.add_data_validation(diff_dv)
    diff_dv.add(f"J2:J1000")

    # 添加说明 sheet
    ws2 = wb.create_sheet(title="填写说明")
    notes = [
        ["字段", "说明", "备注"],
        ["题型", "必填。单选题 / 多选题 / 判断题 / 填空题 / 问答题", "请从下拉列表中选择"],
        ["题干", "必填。试题的题目内容", "与已有试题重复时自动跳过"],
        ["选项A~D", "单选题和多选题必填，其他题型留空", "依次填入选项内容"],
        ["正确答案", "单选题填单个字母，多选题用逗号分隔（如 A,C）", "判断题填 true/false，填空题填答案文本"],
        ["关键词", "问答题的关键词，多个用逗号分隔", "用于自动评分"],
        ["标签", "试题标签，多个用逗号分隔", "如 Python,基础"],
        ["难度", "可选。easy / medium / hard", "默认为 medium"],
        ["", "", ""],
        ["提示", "导入前请先在系统界面选择目标题库分类", "导入的试题将归属到所选分类"],
    ]
    for row_idx, row in enumerate(notes, start=1):
        for col_idx, val in enumerate(row, start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 1:
                cell.font = Font(name="微软雅黑", bold=True, size=11, color="1B365D")
            else:
                cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = thin_border

    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 50
    ws2.column_dimensions["C"].width = 30

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=questions_template.xlsx"},
    )


def _format_answer(q_type: str, answer: dict) -> str:
    """格式化答案用于 Excel 导出"""
    if q_type == "single_choice":
        return str(answer.get("correct", ""))
    elif q_type == "multi_choice":
        correct = answer.get("correct", [])
        return ",".join(correct) if isinstance(correct, list) else str(correct)
    elif q_type == "true_false":
        return "true" if answer.get("correct") else "false"
    elif q_type == "fill_blank":
        return str(answer.get("correct", ""))
    elif q_type == "essay":
        return ""  # 问答题答案不导出
    return ""


@router.post("/questions/import")
async def import_questions(
    file: UploadFile = File(...),
    category_id: str | None = Form(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    """从 Excel 导入题库"""
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 或 .xls 格式")

    contents = await file.read()
    wb = load_workbook(io.BytesIO(contents))
    ws = wb.active

    created = 0
    skipped = 0
    errors = []

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0] or not row[1]:  # 题型和题干必填
            continue

        try:
            q_type_cn = str(row[0]).strip()
            q_type_en = {v: k for k, v in TYPE_MAP.items()}.get(q_type_cn)
            if not q_type_en:
                errors.append(f"第{i}行: 无效题型 '{q_type_cn}'")
                continue

            stem = str(row[1]).strip()

            # 重复检测
            existing = db.query(Question).filter(Question.stem == stem).first()
            if existing:
                skipped += 1
                continue

            # 选项
            options = None
            if q_type_en in ("single_choice", "multi_choice"):
                options = {}
                option_labels = ["A", "B", "C", "D"]
                for j, label in enumerate(option_labels):
                    if row[2 + j] and str(row[2 + j]).strip():
                        options[label] = str(row[2 + j]).strip()

            # 答案
            answer_str = str(row[6]).strip() if row[6] else ""
            answer = _parse_answer(q_type_en, answer_str)

            # 关键词
            keywords = None
            if row[7] and str(row[7]).strip():
                import re
                keywords = [k.strip() for k in re.split(r"[,，]", str(row[7])) if k.strip()]

            # 标签
            tags = None
            if row[8] and str(row[8]).strip():
                tags = [t.strip() for t in str(row[8]).split(",") if t.strip()]

            # 难度
            difficulty = str(row[9]).strip() if row[9] else "medium"
            if difficulty not in ("easy", "medium", "hard"):
                difficulty = "medium"

            question = Question(
                id=str(uuid.uuid4()),
                type=q_type_en,
                stem=stem,
                options=options,
                answer=answer,
                keywords=keywords,
                tags=tags,
                difficulty=difficulty,
                category_id=category_id,
                created_by=current_user.id,
            )
            db.add(question)
            created += 1

        except Exception as e:
            errors.append(f"第{i}行: {str(e)}")

    db.commit()

    write_log(db=db, operator=current_user.name, operator_id=current_user.id, category="question", action="import", resource_type="试题", resource_name="", description=f"{current_user.name} 批量导入了 {created} 道试题", level="info")
    return {"created": created, "skipped": skipped, "errors": errors}


# ========== AI 草稿批量入库 ==========

@router.post("/questions/bulk-create")
def bulk_create_questions(
    data: BulkCreateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    """批量创建试题（AI 生成草稿人工审核后的入库通道）。

    与 Excel 导入同款返回形状 {created, skipped, errors}；逐条再过一遍
    validate_item 双保险（不信任前端数据），归一化题干去重，统一打「AI生成」标签。
    """
    created = 0
    skipped = 0
    errors = []
    seen_keys: set[str] = set()
    existing_keys = {normalize_stem(stem) for (stem,) in db.query(Question.stem).all()}

    for i, item in enumerate(data.items, start=1):
        draft, ws = validate_item(item.model_dump(), expected_type=item.type)
        if draft is None:
            errors.append(f"第{i}题：{ws[-1] if ws else '校验未通过'}")
            continue
        key = normalize_stem(draft["stem"])
        if key in existing_keys or key in seen_keys:
            skipped += 1
            continue
        seen_keys.add(key)
        # 去重保序追加 AI 标签
        tags = list(dict.fromkeys([*(item.tags or []), AI_TAG]))
        question = Question(
            id=str(uuid.uuid4()),
            type=draft["type"],
            stem=draft["stem"],
            options=draft["options"],
            answer=draft["answer"],
            keywords=draft["keywords"],
            difficulty=draft["difficulty"],
            tags=tags,
            section=item.section,
            category_id=item.category_id,
            created_by=current_user.id,
        )
        db.add(question)
        created += 1

    db.commit()

    write_log(
        db=db, operator=current_user.name, operator_id=current_user.id,
        category="question", action="bulk_create", resource_type="试题", resource_name="",
        description=f"{current_user.name} 批量入库了 {created} 道试题（跳过重复 {skipped}，失败 {len(errors)}）",
        level="info",
    )
    return {"created": created, "skipped": skipped, "errors": errors}


# ========== 单题操作 ==========

@router.get("/questions/{question_id}", response_model=QuestionOut)
def get_question(
    question_id: str,
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    question = db.query(Question).filter(Question.id == question_id).first()
    if not question:
        raise HTTPException(status_code=404, detail="试题不存在")
    return question


@router.put("/questions/{question_id}", response_model=QuestionOut)
def update_question(
    question_id: str,
    data: QuestionUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    question = db.query(Question).filter(Question.id == question_id).first()
    if not question:
        raise HTTPException(status_code=404, detail="试题不存在")
    for field, value in data.model_dump(exclude_unset=True).items():
        setattr(question, field, value)
    db.commit()
    db.refresh(question)
    stem = question.stem
    write_log(db=db, operator=current_user.name, operator_id=current_user.id, category="question", action="update", resource_type="试题", resource_name=stem[:50] if stem else "", description=f"{current_user.name} 编辑了试题「{stem[:50]}」", level="info")
    return question


@router.post("/questions/{question_id}/clone", response_model=QuestionOut, status_code=201)
def clone_question(
    question_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    """复制试题"""
    original = db.query(Question).filter(Question.id == question_id).first()
    if not original:
        raise HTTPException(status_code=404, detail="试题不存在")
    question = Question(
        id=str(uuid.uuid4()),
        type=original.type,
        stem=original.stem,
        options=original.options,
        answer=original.answer,
        keywords=original.keywords,
        difficulty=original.difficulty,
        tags=original.tags,
        section=original.section,
        category_id=original.category_id,
        created_by=current_user.id,
    )
    db.add(question)
    db.commit()
    db.refresh(question)
    return question


@router.delete("/questions/{question_id}")
def delete_question(
    question_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    question = db.query(Question).filter(Question.id == question_id).first()
    if not question:
        raise HTTPException(status_code=404, detail="试题不存在")

    refs = _check_references(db, [question_id])
    if refs:
        ref_names = [f"{r['source_type']}: {r['source_name']}" for r in refs]
        raise HTTPException(
            status_code=409,
            detail=f"试题被以下资源引用，无法删除: {', '.join(ref_names)}",
        )

    db.delete(question)
    db.commit()
    stem = question.stem
    write_log(db=db, operator=current_user.name, operator_id=current_user.id, category="question", action="delete", resource_type="试题", resource_name=stem[:50] if stem else "", description=f"{current_user.name} 删除了试题「{stem[:50]}」", level="warning")
    return {"ok": True}


def _check_references(db: Session, question_ids: list[str]) -> list[dict]:
    """检查试题是否被试卷引用，返回引用信息列表"""
    refs: list[dict] = []

    # 检查固定试卷 (Paper.mandatory_questions)
    papers = db.query(Paper).all()
    for paper in papers:
        if paper.mandatory_questions:
            for mq in paper.mandatory_questions:
                if mq.get("question_id") in question_ids:
                    refs.append({
                        "question_id": mq["question_id"],
                        "source_type": "paper",
                        "source_id": paper.id,
                        "source_name": paper.title,
                    })

    # 检查考试实例 (ExamPaper.questions)
    exam_papers = db.query(ExamPaper).all()
    for ep in exam_papers:
        if ep.questions:
            for q in ep.questions:
                if q.get("question_id") in question_ids:
                    refs.append({
                        "question_id": q["question_id"],
                        "source_type": "exam",
                        "source_id": ep.exam_id,
                        "source_name": ep.exam_id,
                    })

    return refs


def _parse_answer(q_type: str, answer_str: str) -> dict:
    """解析 Excel 中的答案字符串"""
    if q_type == "single_choice":
        return {"correct": answer_str}
    elif q_type == "multi_choice":
        return {"correct": [a.strip() for a in answer_str.split(",") if a.strip()]}
    elif q_type == "true_false":
        return {"correct": answer_str.lower() == "true"}
    elif q_type == "fill_blank":
        return {"correct": answer_str}
    elif q_type == "essay":
        return {"reference": answer_str}
    return {"correct": answer_str}